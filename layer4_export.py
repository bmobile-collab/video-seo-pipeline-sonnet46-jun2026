#!/usr/bin/env python3
"""
Amazon Video Pipeline — Layer 4: Excel Export
==============================================
Takes analysis.json + transcript.json from Layer 3
Exports a clean, formatted Excel report with 6 tabs:
  1. Summary       — scores, product info, key stats
  2. Listing Copy  — title options, bullets, backend keywords
  3. Keywords      — buyer phrases ranked by value
  4. Hooks & Ads   — hooks used + ad hooks to steal/beat
  5. Gaps          — competitor weaknesses + your opportunities
  6. Raw Transcript — full transcript with timestamps

Usage:
  python layer4_export.py <video_id>
  python layer4_export.py --batch

Install:
  pip install openpyxl
"""

import sys
import json
import argparse
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter

DOWNLOADS_DIR = Path("downloads")

# Colors
DARK_BLUE   = "1B3A5C"
MID_BLUE    = "2E75B6"
LIGHT_BLUE  = "BDD7EE"
ORANGE      = "E36B1A"
LIGHT_ORANGE= "FCE4D6"
GREEN       = "375623"
LIGHT_GREEN = "E2EFDA"
YELLOW      = "FFF2CC"
WHITE       = "FFFFFF"
LIGHT_GRAY  = "F2F2F2"
MID_GRAY    = "D9D9D9"


def hdr(ws, row, col, value, bg=DARK_BLUE, fg=WHITE, size=11, bold=True, wrap=False):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(name="Arial", bold=bold, color=fg, size=size)
    cell.fill = PatternFill("solid", start_color=bg)
    cell.alignment = Alignment(horizontal="left", vertical="center",
                               wrap_text=wrap)
    return cell


def cell(ws, row, col, value, bg=None, fg="000000", bold=False,
         wrap=True, size=10, align="left"):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name="Arial", bold=bold, color=fg, size=size)
    if bg:
        c.fill = PatternFill("solid", start_color=bg)
    c.alignment = Alignment(horizontal=align, vertical="top", wrap_text=wrap)
    return c


def border_row(ws, row, cols):
    thin = Side(style="thin", color=MID_GRAY)
    for col in range(1, cols + 1):
        ws.cell(row=row, column=col).border = Border(bottom=thin)


def section_title(ws, row, col, text, colspan=1):
    c = ws.cell(row=row, column=col, value=text)
    c.font = Font(name="Arial", bold=True, color=DARK_BLUE, size=11)
    c.fill = PatternFill("solid", start_color=LIGHT_BLUE)
    c.alignment = Alignment(horizontal="left", vertical="center")
    if colspan > 1:
        ws.merge_cells(start_row=row, start_column=col,
                       end_row=row, end_column=col + colspan - 1)
    return c


def build_summary(wb, analysis, meta, video_id):
    ws = wb.create_sheet("Summary")
    ws.sheet_view.showGridLines = False

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 60
    ws.row_dimensions[1].height = 36

    # Title bar
    ws.merge_cells("A1:B1")
    hdr(ws, 1, 1, "🎯  AMAZON VIDEO INTELLIGENCE REPORT", DARK_BLUE, WHITE, 14)

    ws.merge_cells("A2:B2")
    hdr(ws, 2, 1, f"Video ID: {video_id}  |  Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
        MID_BLUE, WHITE, 10, bold=False)

    r = 4
    section_title(ws, r, 1, "PRODUCT OVERVIEW", 2); r += 1
    fields = [
        ("Product", analysis.get("product_summary", "")),
        ("Target Audience", analysis.get("target_audience", "")),
        ("Source URL", meta.get("url", "")),
    ]
    for label, value in fields:
        cell(ws, r, 1, label, LIGHT_GRAY, bold=True)
        cell(ws, r, 2, value, wrap=True)
        border_row(ws, r, 2)
        r += 1

    r += 1
    section_title(ws, r, 1, "COMPETITOR SCORE CARD", 2); r += 1
    hdr(ws, r, 1, "Metric", LIGHT_BLUE, DARK_BLUE, 10)
    hdr(ws, r, 2, "Score", LIGHT_BLUE, DARK_BLUE, 10); r += 1

    score = analysis.get("overall_score", {})
    score_rows = [
        ("Listing Strength", f"{score.get('listing_strength', '?')} / 10"),
        ("Hook Quality",     f"{score.get('hook_quality', '?')} / 10"),
        ("Keyword Density",  f"{score.get('keyword_density', '?')} / 10"),
        ("Assessment",       score.get("notes", "")),
    ]
    for i, (label, value) in enumerate(score_rows):
        bg = LIGHT_GRAY if i % 2 == 0 else WHITE
        cell(ws, r, 1, label, bg, bold=True)
        cell(ws, r, 2, value, bg)
        border_row(ws, r, 2)
        r += 1

    r += 1
    section_title(ws, r, 1, "PIPELINE STATS", 2); r += 1
    trans = meta.get("transcription", {})
    stats = [
        ("Words Transcribed",  trans.get("word_count", "—")),
        ("Duration",           f"{trans.get('duration_seconds', 0):.0f} seconds"),
        ("Whisper Model",      trans.get("model", "—")),
        ("Analysis Model",     meta.get("analysis", {}).get("model", "—")),
        ("Hooks Found",        len(analysis.get("hooks", []))),
        ("Gaps Identified",    len(analysis.get("competitor_gaps", []))),
        ("Buyer Phrases",      len(analysis.get("buyer_language", []))),
    ]
    for i, (label, value) in enumerate(stats):
        bg = LIGHT_GRAY if i % 2 == 0 else WHITE
        cell(ws, r, 1, label, bg, bold=True)
        cell(ws, r, 2, str(value), bg)
        border_row(ws, r, 2)
        r += 1

    ws.row_dimensions[r].height = 6
    return ws


def build_listing_copy(wb, analysis):
    ws = wb.create_sheet("Listing Copy")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 100

    ws.merge_cells("A1:B1")
    hdr(ws, 1, 1, "📝  READY-TO-PASTE LISTING COPY", DARK_BLUE, WHITE, 13)

    r = 3
    section_title(ws, r, 1, "TITLE OPTIONS", 2); r += 1
    hdr(ws, r, 1, "#", LIGHT_BLUE, DARK_BLUE, 10)
    hdr(ws, r, 2, "Title", LIGHT_BLUE, DARK_BLUE, 10); r += 1
    for i, title in enumerate(analysis.get("listing_title_suggestions", []), 1):
        bg = LIGHT_GRAY if i % 2 == 0 else WHITE
        cell(ws, r, 1, str(i), bg, align="center")
        cell(ws, r, 2, title, bg, wrap=True)
        ws.row_dimensions[r].height = 30
        border_row(ws, r, 2)
        r += 1

    r += 1
    section_title(ws, r, 1, "BULLET POINTS (copy-paste order)", 2); r += 1
    hdr(ws, r, 1, "#", LIGHT_BLUE, DARK_BLUE, 10)
    hdr(ws, r, 2, "Bullet", LIGHT_BLUE, DARK_BLUE, 10); r += 1
    for i, bullet in enumerate(analysis.get("bullet_points", []), 1):
        bg = LIGHT_GREEN if i % 2 == 0 else WHITE
        cell(ws, r, 1, str(i), bg, align="center")
        cell(ws, r, 2, bullet, bg, wrap=True)
        ws.row_dimensions[r].height = 45
        border_row(ws, r, 2)
        r += 1

    r += 1
    section_title(ws, r, 1, "BACKEND KEYWORDS", 2); r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    c = ws.cell(row=r, column=1, value=analysis.get("backend_keywords", ""))
    c.font = Font(name="Arial", size=10)
    c.fill = PatternFill("solid", start_color=YELLOW)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 60

    return ws


def build_keywords(wb, analysis):
    ws = wb.create_sheet("Keywords")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 25

    ws.merge_cells("A1:C1")
    hdr(ws, 1, 1, "🔑  BUYER LANGUAGE & KEYWORDS", DARK_BLUE, WHITE, 13)

    r = 3
    hdr(ws, r, 1, "Phrase", MID_BLUE, WHITE, 10)
    hdr(ws, r, 2, "Value", MID_BLUE, WHITE, 10)
    hdr(ws, r, 3, "Use In", MID_BLUE, WHITE, 10); r += 1

    color_map = {"high": LIGHT_GREEN, "medium": YELLOW, "low": LIGHT_GRAY}
    for item in sorted(analysis.get("buyer_language", []),
                       key=lambda x: {"high": 0, "medium": 1, "low": 2}.get(x.get("keyword_value", "low"), 2)):
        val = item.get("keyword_value", "").lower()
        bg = color_map.get(val, WHITE)
        cell(ws, r, 1, item.get("phrase", ""), bg, wrap=True)
        cell(ws, r, 2, val.upper(), bg, bold=True, align="center")
        cell(ws, r, 3, item.get("use_in", ""), bg)
        border_row(ws, r, 3)
        r += 1

    r += 1
    section_title(ws, r, 1, "BENEFITS CLAIMED BY COMPETITOR", 3); r += 1
    hdr(ws, r, 1, "Phrase", LIGHT_BLUE, DARK_BLUE, 10)
    hdr(ws, r, 2, "Type", LIGHT_BLUE, DARK_BLUE, 10)
    hdr(ws, r, 3, "Strength", LIGHT_BLUE, DARK_BLUE, 10); r += 1
    for i, b in enumerate(analysis.get("benefits_claimed", [])):
        bg = LIGHT_GRAY if i % 2 == 0 else WHITE
        cell(ws, r, 1, b.get("phrase", ""), bg, wrap=True)
        cell(ws, r, 2, b.get("benefit_type", ""), bg)
        cell(ws, r, 3, b.get("strength", ""), bg)
        border_row(ws, r, 3)
        r += 1

    return ws


def build_hooks(wb, analysis):
    ws = wb.create_sheet("Hooks & Ad Copy")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 20

    ws.merge_cells("A1:C1")
    hdr(ws, 1, 1, "🎣  HOOKS & AD COPY INTELLIGENCE", DARK_BLUE, WHITE, 13)

    r = 3
    section_title(ws, r, 1, "HOOKS COMPETITOR USED", 3); r += 1
    hdr(ws, r, 1, "Type", MID_BLUE, WHITE, 10)
    hdr(ws, r, 2, "Hook Text", MID_BLUE, WHITE, 10)
    hdr(ws, r, 3, "Timestamp", MID_BLUE, WHITE, 10); r += 1
    for i, h in enumerate(analysis.get("hooks", [])):
        bg = LIGHT_ORANGE if i % 2 == 0 else WHITE
        cell(ws, r, 1, h.get("type", "").upper(), bg, bold=True)
        cell(ws, r, 2, h.get("text", ""), bg, wrap=True)
        cell(ws, r, 3, h.get("timestamp_approx", ""), bg)
        ws.row_dimensions[r].height = 35
        border_row(ws, r, 3)
        r += 1

    r += 1
    section_title(ws, r, 1, "AD HOOKS TO STEAL / BEAT", 3); r += 1
    hdr(ws, r, 1, "#", LIGHT_BLUE, DARK_BLUE, 10)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    hdr(ws, r, 2, "Hook", LIGHT_BLUE, DARK_BLUE, 10); r += 1
    for i, hook in enumerate(analysis.get("ad_hooks", []), 1):
        bg = LIGHT_GREEN if i % 2 == 0 else WHITE
        cell(ws, r, 1, str(i), bg, align="center", bold=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        cell(ws, r, 2, hook, bg, wrap=True)
        ws.row_dimensions[r].height = 40
        border_row(ws, r, 3)
        r += 1

    return ws


def build_gaps(wb, analysis):
    ws = wb.create_sheet("Competitor Gaps")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 50
    ws.column_dimensions["B"].width = 55

    ws.merge_cells("A1:B1")
    hdr(ws, 1, 1, "🚨  COMPETITOR GAPS — YOUR OPPORTUNITIES", DARK_BLUE, WHITE, 13)

    r = 3
    hdr(ws, r, 1, "Gap / Weakness", ORANGE, WHITE, 10)
    hdr(ws, r, 2, "Your Opportunity", GREEN, WHITE, 10); r += 1

    for i, g in enumerate(analysis.get("competitor_gaps", [])):
        bg_l = LIGHT_ORANGE if i % 2 == 0 else WHITE
        bg_r = LIGHT_GREEN if i % 2 == 0 else WHITE
        cell(ws, r, 1, g.get("gap", ""), bg_l, wrap=True)
        cell(ws, r, 2, g.get("your_opportunity", ""), bg_r, wrap=True)
        ws.row_dimensions[r].height = 50
        border_row(ws, r, 2)
        r += 1

    r += 1
    section_title(ws, r, 1, "PAIN POINTS THEY TARGETED", 2); r += 1
    hdr(ws, r, 1, "Pain Point", LIGHT_BLUE, DARK_BLUE, 10)
    hdr(ws, r, 2, "Your Counter-Opportunity", LIGHT_BLUE, DARK_BLUE, 10); r += 1
    for i, p in enumerate(analysis.get("pain_points", [])):
        bg = LIGHT_GRAY if i % 2 == 0 else WHITE
        cell(ws, r, 1, p.get("pain", ""), bg, wrap=True)
        cell(ws, r, 2, p.get("opportunity", ""), bg, wrap=True)
        ws.row_dimensions[r].height = 45
        border_row(ws, r, 2)
        r += 1

    return ws


def build_transcript(wb, transcript_data):
    ws = wb.create_sheet("Raw Transcript")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 90

    ws.merge_cells("A1:C1")
    hdr(ws, 1, 1, "📄  RAW TRANSCRIPT WITH TIMESTAMPS", DARK_BLUE, WHITE, 13)

    r = 3
    hdr(ws, r, 1, "Start (s)", MID_BLUE, WHITE, 10)
    hdr(ws, r, 2, "End (s)", MID_BLUE, WHITE, 10)
    hdr(ws, r, 3, "Text", MID_BLUE, WHITE, 10); r += 1

    segments = transcript_data.get("segments", [])
    if segments:
        for i, seg in enumerate(segments):
            bg = LIGHT_GRAY if i % 2 == 0 else WHITE
            cell(ws, r, 1, seg.get("start", ""), bg, align="center")
            cell(ws, r, 2, seg.get("end", ""), bg, align="center")
            cell(ws, r, 3, seg.get("text", ""), bg, wrap=True)
            ws.row_dimensions[r].height = 30
            border_row(ws, r, 3)
            r += 1
    else:
        full = transcript_data.get("full_text", "")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        c = ws.cell(row=r, column=1, value=full)
        c.font = Font(name="Arial", size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = max(100, len(full) // 3)

    return ws


def export(video_id: str) -> Path:
    out_dir = DOWNLOADS_DIR / video_id

    analysis_path = out_dir / "analysis.json"
    transcript_path = out_dir / "transcript.json"
    meta_path = out_dir / "meta.json"

    if not analysis_path.exists():
        print(f"  ❌  analysis.json not found. Run layer3_analyze.py first.")
        return None

    analysis = json.loads(analysis_path.read_text(encoding="utf-8"))
    transcript_data = json.loads(transcript_path.read_text(encoding="utf-8")) if transcript_path.exists() else {}
    meta = json.loads(meta_path.read_text(encoding="utf-8")) if meta_path.exists() else {}

    print(f"\n{'='*55}")
    print(f"  Video ID : {video_id}")
    print(f"  Building Excel report...")
    print(f"{'='*55}")

    wb = Workbook()
    wb.remove(wb.active)

    build_summary(wb, analysis, meta, video_id)
    print(f"  ✅  Tab 1: Summary")
    build_listing_copy(wb, analysis)
    print(f"  ✅  Tab 2: Listing Copy")
    build_keywords(wb, analysis)
    print(f"  ✅  Tab 3: Keywords")
    build_hooks(wb, analysis)
    print(f"  ✅  Tab 4: Hooks & Ad Copy")
    build_gaps(wb, analysis)
    print(f"  ✅  Tab 5: Competitor Gaps")
    build_transcript(wb, transcript_data)
    print(f"  ✅  Tab 6: Raw Transcript")

    ts = datetime.now().strftime("%Y%m%d_%H%M")
    xlsx_path = out_dir / f"intelligence_report_{ts}.xlsx"
    wb.save(str(xlsx_path))

    size_kb = xlsx_path.stat().st_size / 1024
    print(f"\n  📊  Saved: {xlsx_path}")
    print(f"      Size : {size_kb:.1f} KB")
    print(f"\n  ✅  Layer 4 complete!")
    print(f"      Open: {xlsx_path}")

    if meta_path.exists():
        meta["export"] = {
            "status": "complete",
            "exported_at": datetime.utcnow().isoformat() + "Z",
            "file": str(xlsx_path)
        }
        meta["next_step"] = "pipeline complete"
        meta_path.write_text(json.dumps(meta, indent=2))

    return xlsx_path


def main():
    parser = argparse.ArgumentParser(description="Amazon Video Pipeline — Layer 4: Excel Export")
    parser.add_argument("video_id", nargs="?", help="Video ID")
    parser.add_argument("--batch", action="store_true", help="Export all analyzed videos")
    args = parser.parse_args()

    if args.batch:
        if not DOWNLOADS_DIR.exists():
            print("  ❌  No downloads folder found.")
            return
        pending = [
            f.name for f in DOWNLOADS_DIR.iterdir()
            if f.is_dir() and (f / "analysis.json").exists()
        ]
        if not pending:
            print("  ✅  Nothing to export.")
            return
        print(f"  📋  {len(pending)} video(s) to export")
        for i, vid_id in enumerate(pending, 1):
            print(f"\n  [{i}/{len(pending)}]")
            export(vid_id)
    elif args.video_id:
        export(args.video_id)
    else:
        parser.print_help()
        print("\n  Example:")
        print("    python layer4_export.py 0287b3106d634bd189057a96fc83101b")


if __name__ == "__main__":
    main()
